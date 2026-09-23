"""Bộ kết xuất báo cáo toàn văn: cùng một cây nội dung -> .docx và .pdf.

Khối nội dung (xem scripts/report/content.py) là tuple:
    ("h1", text) | ("h2", text) | ("h3", text)
    ("p", text) | ("p_noindent", text) | ("quote", text)
    ("bullet", [item, ...]) | ("num", [item, ...])
    ("table", caption, headers, rows, aligns)
    ("figure", caption, relative_png_path)
    ("pagebreak",)
Cú pháp inline: **đậm**, *nghiêng*, `mã`.
"""
from __future__ import annotations

import re
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]

CM = 28.3464567  # 1 cm tinh theo point
PAGE_W, PAGE_H = 21.0 * CM, 29.7 * CM
M_LEFT, M_RIGHT, M_TOP, M_BOTTOM = 3.0 * CM, 2.0 * CM, 2.0 * CM, 2.0 * CM
TEXT_W = PAGE_W - M_LEFT - M_RIGHT

BODY_PT = 13
LEADING = 1.4
INDENT = 1.0 * CM

FONT_DIR = Path("/System/Library/Fonts/Supplemental")
FONT_FILES = {
    "": FONT_DIR / "Times New Roman.ttf",
    "b": FONT_DIR / "Times New Roman Bold.ttf",
    "i": FONT_DIR / "Times New Roman Italic.ttf",
    "bi": FONT_DIR / "Times New Roman Bold Italic.ttf",
}

_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def split_inline(text: str):
    """Tách chuỗi thành [(text, bold, italic, mono), ...]."""
    out = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False, False))
        elif part.startswith("*") and part.endswith("*"):
            out.append((part[1:-1], False, True, False))
        elif part.startswith("`") and part.endswith("`"):
            out.append((part[1:-1], False, False, True))
        else:
            out.append((part, False, False, False))
    return out


def strip_inline(text: str) -> str:
    return re.sub(r"[*`]", "", text)


# Khổ hiển thị tối đa của hình trong khung soạn thảo
FIG_MAX_W = TEXT_W * 0.86
FIG_MAX_H = 0.33 * (PAGE_H - M_TOP - M_BOTTOM)
TARGET_DPI = 300
LOGO_W = 3.6 * CM  # bề rộng logo trên trang bìa

_CACHE_DIR = Path(tempfile.gettempdir()) / "nlp_report_figs"


def fig_size(path: Path, max_w: float = FIG_MAX_W, max_h: float = FIG_MAX_H):
    """Kích thước hiển thị (point) giữ đúng tỷ lệ ảnh."""
    from PIL import Image

    with Image.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h, 1.0)
    return w * scale, h * scale


def fig_asset(path: Path):
    """Trả về (đường dẫn ảnh, rộng, cao tính theo point).

    Ảnh gốc 300 dpi thường rộng hơn nhiều so với khổ in thực tế, làm tệp kết
    quả phình lên vô ích. Hàm này tạo bản sao đã thu nhỏ về đúng 300 dpi tại
    khổ hiển thị — chất lượng in không đổi, dung lượng giảm mạnh.
    """
    from PIL import Image

    disp_w, disp_h = fig_size(path)
    need_px = int(round(disp_w / 72.0 * TARGET_DPI))

    with Image.open(path) as im:
        if im.width <= need_px * 1.05:
            return path, disp_w, disp_h
        _CACHE_DIR.mkdir(parents=True, exist_ok=True)
        cached = _CACHE_DIR / f"{path.stem}_{need_px}.png"
        if not cached.exists():
            ratio = need_px / im.width
            resized = im.convert("RGB").resize(
                (need_px, max(1, int(round(im.height * ratio)))), Image.LANCZOS)
            resized.save(cached, format="PNG", optimize=True, dpi=(TARGET_DPI, TARGET_DPI))
    return cached, disp_w, disp_h


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def build_docx(blocks, out_path: Path, meta: dict):
    import docx
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = docx.Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)

    def set_font(style, size=BODY_PT, bold=False, italic=False):
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), "Times New Roman")

    normal = doc.styles["Normal"]
    set_font(normal)
    pf = normal.paragraph_format
    pf.line_spacing = LEADING
    pf.space_after = Pt(6)
    pf.first_line_indent = Cm(1)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, bold, italic, before, after in (
        ("Heading 1", 14, True, False, 18, 10),
        ("Heading 2", 13, True, False, 12, 6),
        ("Heading 3", 13, True, True, 10, 6),
    ):
        st = doc.styles[name]
        set_font(st, size, bold, italic)
        st.paragraph_format.line_spacing = LEADING
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def add_runs(par, text):
        for chunk, bold, italic, mono in split_inline(text):
            run = par.add_run(chunk)
            run.bold = bold
            run.italic = italic
            if mono:
                run.font.name = "Consolas"
                run.font.size = Pt(BODY_PT - 1)
                rpr = run._element.get_or_add_rPr()
                rf = OxmlElement("w:rFonts")
                for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                    rf.set(qn(attr), "Consolas")
                rpr.append(rf)
        return par

    def plain(text, *, indent=True, align=None, size=None, italic=False,
              space_before=0, space_after=6, bold=False):
        par = doc.add_paragraph()
        add_runs(par, text)
        pfm = par.paragraph_format
        pfm.first_line_indent = Cm(1 if indent else 0)
        pfm.space_before = Pt(space_before)
        pfm.space_after = Pt(space_after)
        pfm.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
        if size or italic or bold:
            for run in par.runs:
                if size:
                    run.font.size = Pt(size)
                if italic:
                    run.italic = True
                if bold:
                    run.bold = True
        return par

    def field(par, instr):
        r1 = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), "begin")
        r1.append(fc)
        r2 = OxmlElement("w:r")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r2.append(it)
        r3 = OxmlElement("w:r")
        fs = OxmlElement("w:fldChar")
        fs.set(qn("w:fldCharType"), "separate")
        r3.append(fs)
        r4 = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r4.append(t)
        r5 = OxmlElement("w:r")
        fe = OxmlElement("w:fldChar")
        fe.set(qn("w:fldCharType"), "end")
        r5.append(fe)
        for el in (r1, r2, r3, r4, r5):
            par._p.append(el)

    # ---- footer: số trang căn giữa ----
    footer_par = sec.footer.paragraphs[0]
    footer_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_par.paragraph_format.first_line_indent = Cm(0)
    field(footer_par, " PAGE ")
    for run in footer_par.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    # ---- viền trang, chỉ áp cho trang bìa ----
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    pg_borders.set(qn("w:display"), "firstPage")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "triple")
        edge.set(qn("w:sz"), "18")
        edge.set(qn("w:space"), "24")
        edge.set(qn("w:color"), "000000")
        pg_borders.append(edge)
    # Thứ tự phần tử con của w:sectPr do schema quy định: pgBorders phải đứng
    # ngay sau pgMar, không được append vào cuối.
    sec._sectPr.find(qn("w:pgMar")).addnext(pg_borders)
    sec.different_first_page_header_footer = True  # trang bìa không đánh số

    # ---- trang bìa ----
    def cover_line(text, size, bold=False, space_after=6, caps=False,
                   align=WD_ALIGN_PARAGRAPH.CENTER):
        par = doc.add_paragraph()
        run = par.add_run(text.upper() if caps else text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.bold = bold
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.line_spacing = 1.15
        par.paragraph_format.alignment = align
        return par

    LEFT = WD_ALIGN_PARAGRAPH.LEFT

    cover_line(meta["org"], 13, True, 2, caps=True)
    cover_line(meta["school"], 13, True, 2, caps=True)
    cover_line(meta["faculty"], 13, True, 2, caps=True)
    cover_line("-o0o-", 13, True, 10)

    logo_par = doc.add_paragraph()
    logo_par.paragraph_format.first_line_indent = Cm(0)
    logo_par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_par.paragraph_format.space_after = Pt(14)
    logo_w, _ = fig_size(PROJECT_ROOT / meta["logo"], LOGO_W, LOGO_W)
    logo_par.add_run().add_picture(str(PROJECT_ROOT / meta["logo"]), width=Pt(logo_w))

    cover_line(meta["doc_kind"], 20, True, 6, caps=True)
    cover_line(meta["course"], 15, True, 16, caps=True)
    cover_line(meta["topic_label"], 14, True, 4, caps=True)
    cover_line(f'"{meta["title"].upper()}"', 15, True, 22)

    cover_line(meta["advisor_label"], 13, True, 2, caps=True, align=LEFT)
    cover_line(meta["advisor"], 13, True, 12, align=LEFT)
    cover_line(meta["team_label"], 13, True, 2, caps=True, align=LEFT)
    cover_line(meta["team_name"], 13, True, 2, align=LEFT)
    for text, bold in meta["members"]:
        cover_line(text, 13, bold, 2, align=LEFT)

    cover_line("", 13, False, 20)
    cover_line(meta["place_date"], 14, True, 0, caps=True)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- nội dung ----
    for block in blocks:
        kind = block[0]

        if kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        elif kind == "toc":
            par = doc.add_paragraph()
            par.paragraph_format.first_line_indent = Cm(0)
            field(par, r' TOC \o "1-3" \h \z \u ')
            plain("(Mở bằng Microsoft Word và nhấn Ctrl+A rồi F9 để cập nhật mục lục.)",
                  indent=False, italic=True, size=11)

        elif kind in ("h1", "h2", "h3"):
            level = int(kind[1])
            par = doc.add_heading(level=level)
            add_runs(par, block[1])
            for run in par.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.name = "Times New Roman"
                run.font.size = Pt(14 if level == 1 else 13)
                run.bold = True
                run.italic = level == 3

        elif kind == "p":
            plain(block[1])

        elif kind == "p_noindent":
            plain(block[1], indent=False)

        elif kind == "lines":
            for line in block[1]:
                plain(line, indent=False, space_after=3)

        elif kind == "quote":
            par = plain(block[1], indent=False, italic=True, space_before=4, space_after=8)
            par.paragraph_format.left_indent = Cm(1.0)

        elif kind in ("bullet", "num"):
            style = "List Bullet" if kind == "bullet" else "List Number"
            for item in block[1]:
                par = doc.add_paragraph(style=style)
                add_runs(par, item)
                par.paragraph_format.line_spacing = LEADING
                par.paragraph_format.space_after = Pt(3)
                par.paragraph_format.left_indent = Cm(1.0)
                par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                for run in par.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(BODY_PT)

        elif kind == "table":
            _, caption, headers, rows, aligns = block
            cap = plain(caption, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER,
                        size=12, italic=True, space_before=6, space_after=4)
            cap.paragraph_format.keep_with_next = True
            tbl = doc.add_table(rows=1, cols=len(headers))
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = tbl.rows[0].cells
            for i, head in enumerate(headers):
                hdr[i].text = ""
                par = hdr[i].paragraphs[0]
                add_runs(par, head)
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.bold = True
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "E8EDF3")
                hdr[i]._tc.get_or_add_tcPr().append(shd)
            for row in rows:
                cells = tbl.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = ""
                    par = cells[i].paragraphs[0]
                    add_runs(par, str(val))
                    par.paragraph_format.first_line_indent = Cm(0)
                    par.paragraph_format.space_after = Pt(2)
                    al = aligns[i] if i < len(aligns) else "l"
                    par.paragraph_format.alignment = {
                        "l": WD_ALIGN_PARAGRAPH.LEFT,
                        "c": WD_ALIGN_PARAGRAPH.CENTER,
                        "r": WD_ALIGN_PARAGRAPH.RIGHT,
                    }[al]
                    for run in par.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12)
            # lặp lại dòng tiêu đề khi bảng tràn trang
            tr_pr = tbl.rows[0]._tr.get_or_add_trPr()
            hdr_el = OxmlElement("w:tblHeader")
            hdr_el.set(qn("w:val"), "true")
            tr_pr.append(hdr_el)
            plain("", indent=False, space_after=2)

        elif kind == "figure":
            _, caption, rel = block
            path = PROJECT_ROOT / rel
            par = doc.add_paragraph()
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(8)
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.keep_with_next = True
            asset, disp_w, _ = fig_asset(path)
            par.add_run().add_picture(str(asset), width=Pt(disp_w))
            cap = plain(caption, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER,
                        size=12, italic=True, space_after=10)
            cap.paragraph_format.keep_with_next = False

        else:  # pragma: no cover
            raise ValueError(f"Khối không hỗ trợ: {kind}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def build_pdf(blocks, out_path: Path, meta: dict):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Frame, Image,
                                    KeepTogether, ListFlowable, ListItem,
                                    NextPageTemplate, PageBreak, PageTemplate,
                                    Paragraph, Spacer, Table, TableStyle)
    from reportlab.platypus.tableofcontents import TableOfContents

    for suffix, name in (("", "TNR"), ("b", "TNR-Bold"), ("i", "TNR-Italic"), ("bi", "TNR-BoldItalic")):
        pdfmetrics.registerFont(TTFont(name, str(FONT_FILES[suffix])))
    pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNR-Bold",
                                  italic="TNR-Italic", boldItalic="TNR-BoldItalic")

    body = ParagraphStyle("body", fontName="TNR", fontSize=BODY_PT,
                          leading=BODY_PT * LEADING, alignment=TA_JUSTIFY,
                          firstLineIndent=INDENT, spaceAfter=6)
    body_flat = ParagraphStyle("body_flat", parent=body, firstLineIndent=0)
    quote = ParagraphStyle("quote", parent=body_flat, leftIndent=INDENT,
                           fontName="TNR-Italic", spaceBefore=4, spaceAfter=8)
    h1 = ParagraphStyle("h1", fontName="TNR-Bold", fontSize=14, leading=14 * 1.3,
                        alignment=TA_LEFT, spaceBefore=16, spaceAfter=10, firstLineIndent=0)
    h2 = ParagraphStyle("h2", parent=h1, fontSize=13, leading=13 * 1.3,
                        spaceBefore=12, spaceAfter=6)
    h3 = ParagraphStyle("h3", parent=h2, fontName="TNR-BoldItalic",
                        spaceBefore=10, spaceAfter=5)
    caption = ParagraphStyle("caption", fontName="TNR-Italic", fontSize=12,
                             leading=12 * 1.25, alignment=TA_CENTER,
                             spaceBefore=4, spaceAfter=8, firstLineIndent=0)
    cell = ParagraphStyle("cell", fontName="TNR", fontSize=11.5,
                          leading=11.5 * 1.2, firstLineIndent=0)
    cell_c = ParagraphStyle("cell_c", parent=cell, alignment=TA_CENTER)
    cell_r = ParagraphStyle("cell_r", parent=cell, alignment=TA_RIGHT)
    cell_h = ParagraphStyle("cell_h", parent=cell_c, fontName="TNR-Bold")
    cover = ParagraphStyle("cover", fontName="TNR", fontSize=13,
                           leading=13 * 1.4, alignment=TA_CENTER, firstLineIndent=0)

    def esc(text: str) -> str:
        out = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
        for raw, bold, italic, mono in []:
            pass
        return out

    def rich(text: str) -> str:
        parts = []
        for chunk, bold, italic, mono in split_inline(text):
            safe = chunk.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if mono:
                safe = f'<font name="Courier" size="{BODY_PT - 1.5}">{safe}</font>'
            if bold:
                safe = f"<b>{safe}</b>"
            if italic:
                safe = f"<i>{safe}</i>"
            parts.append(safe)
        return "".join(parts)

    story = []

    # --- trang bìa ---
    def cov(text, size, bold=False, gap=6, caps=False, align=TA_CENTER):
        st = ParagraphStyle(f"cov{size}{bold}{align}", parent=cover, fontSize=size,
                            leading=size * 1.3, alignment=align,
                            fontName="TNR-Bold" if bold else "TNR")
        story.append(Paragraph(rich(text.upper() if caps else text), st))
        if gap:
            story.append(Spacer(1, gap))

    cov(meta["org"], 13, True, 2, caps=True)
    cov(meta["school"], 13, True, 2, caps=True)
    cov(meta["faculty"], 13, True, 2, caps=True)
    cov("-o0o-", 13, True, 12)

    logo_path = PROJECT_ROOT / meta["logo"]
    lw, lh = fig_size(logo_path, LOGO_W, LOGO_W)
    story.append(Image(str(logo_path), width=lw, height=lh, hAlign="CENTER"))
    story.append(Spacer(1, 18))

    cov(meta["doc_kind"], 20, True, 8, caps=True)
    cov(meta["course"], 15, True, 20, caps=True)
    cov(meta["topic_label"], 14, True, 6, caps=True)
    cov(f'"{meta["title"].upper()}"', 15, True, 28)

    cov(meta["advisor_label"], 13, True, 3, caps=True, align=TA_LEFT)
    cov(meta["advisor"], 13, True, 16, align=TA_LEFT)
    cov(meta["team_label"], 13, True, 3, caps=True, align=TA_LEFT)
    cov(meta["team_name"], 13, True, 3, align=TA_LEFT)
    for text, bold in meta["members"]:
        cov(text, 13, bold, 3, align=TA_LEFT)

    story.append(Spacer(1, 30))
    cov(meta["place_date"], 14, True, 0, caps=True)
    story.append(PageBreak())

    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle("toc0", fontName="TNR-Bold", fontSize=13, leading=13 * 1.5,
                       leftIndent=0, firstLineIndent=0, spaceBefore=6),
        ParagraphStyle("toc1", fontName="TNR", fontSize=13, leading=13 * 1.45,
                       leftIndent=18, firstLineIndent=0),
        ParagraphStyle("toc2", fontName="TNR-Italic", fontSize=12.5,
                       leading=12.5 * 1.4, leftIndent=36, firstLineIndent=0),
    ]
    toc.dotsMinLevel = 0

    class TocHeading(Paragraph):
        """Paragraph mang metadata để doc template phát tín hiệu TOCEntry."""

        def __init__(self, text, style, level, key):
            super().__init__(text, style)
            self._toc_level = level
            self._toc_key = key

    class ReportDoc(BaseDocTemplate):
        """BaseDocTemplate biết ghi mục lục từ các TocHeading đã vẽ."""

        def afterFlowable(self, flowable):
            level = getattr(flowable, "_toc_level", None)
            if level is None:
                return
            key = flowable._toc_key
            self.canv.bookmarkPage(key)
            self.notify("TOCEntry",
                        (level, strip_inline(flowable.getPlainText()), self.page, key))

    counter = {"n": 0}

    def heading(text, level):
        counter["n"] += 1
        style = (h1, h2, h3)[level]
        return TocHeading(rich(text), style, level, f"h{counter['n']}")

    for block in blocks:
        kind = block[0]

        if kind == "pagebreak":
            story.append(PageBreak())

        elif kind == "toc":
            story.append(toc)

        elif kind in ("h1", "h2", "h3"):
            level = int(kind[1]) - 1
            if level == 0:
                story.append(heading(block[1], level))
            else:
                story.append(CondPageBreak(3 * CM))
                story.append(heading(block[1], level))

        elif kind == "p":
            story.append(Paragraph(rich(block[1]), body))

        elif kind == "p_noindent":
            story.append(Paragraph(rich(block[1]), body_flat))

        elif kind == "lines":
            listing = ParagraphStyle("listing", parent=body_flat, spaceAfter=3)
            for line in block[1]:
                story.append(Paragraph(rich(line), listing))

        elif kind == "quote":
            story.append(Paragraph(rich(block[1]), quote))

        elif kind in ("bullet", "num"):
            items = [ListItem(Paragraph(rich(i), body_flat), leftIndent=INDENT + 14)
                     for i in block[1]]
            story.append(ListFlowable(
                items,
                bulletType="bullet" if kind == "bullet" else "1",
                start="–" if kind == "bullet" else 1,
                leftIndent=INDENT + 14, bulletFontName="TNR", bulletFontSize=BODY_PT,
                spaceAfter=6))

        elif kind == "table":
            _, cap_text, headers, rows, aligns = block
            ncol = len(headers)
            weights = []
            for i in range(ncol):
                longest = max([len(strip_inline(str(headers[i])))] +
                              [len(strip_inline(str(r[i]))) for r in rows] or [1])
                weights.append(max(longest, 4) ** 0.62)
            total = sum(weights)
            widths = [TEXT_W * w / total for w in weights]
            styles_by_align = {"l": cell, "c": cell_c, "r": cell_r}
            data = [[Paragraph(rich(str(h)), cell_h) for h in headers]]
            for row in rows:
                data.append([
                    Paragraph(rich(str(v)),
                              styles_by_align[aligns[i] if i < len(aligns) else "l"])
                    for i, v in enumerate(row)])
            tbl = Table(data, colWidths=widths, repeatRows=1, hAlign="CENTER")
            tbl.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#6b7785")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF3")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(Spacer(1, 4))
            story.append(Paragraph(rich(cap_text), caption))
            story.append(tbl)
            story.append(Spacer(1, 10))

        elif kind == "figure":
            _, cap_text, rel = block
            path = PROJECT_ROOT / rel
            asset, w, h = fig_asset(path)
            story.append(KeepTogether([
                Spacer(1, 4),
                Image(str(asset), width=w, height=h, hAlign="CENTER"),
                Spacer(1, 4),
                Paragraph(rich(cap_text), caption),
            ]))

        else:  # pragma: no cover
            raise ValueError(f"Khối không hỗ trợ: {kind}")

    def on_page(canv, doc_):
        canv.saveState()
        if doc_.page == 1:
            # khung viền ba nét của trang bìa
            for inset, width in ((0.95 * CM, 3.2), (1.22 * CM, 1.0), (1.42 * CM, 1.8)):
                canv.setLineWidth(width)
                canv.rect(inset, inset, PAGE_W - 2 * inset, PAGE_H - 2 * inset)
        else:
            canv.setFont("TNR", 12)
            canv.drawCentredString(PAGE_W / 2, M_BOTTOM - 0.85 * CM, str(doc_.page))
        canv.restoreState()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = ReportDoc(
        str(out_path), pagesize=(PAGE_W, PAGE_H),
        leftMargin=M_LEFT, rightMargin=M_RIGHT,
        topMargin=M_TOP, bottomMargin=M_BOTTOM,
        title=meta["title"], author=meta["author"])
    frame = Frame(M_LEFT, M_BOTTOM, TEXT_W, PAGE_H - M_TOP - M_BOTTOM, id="body")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=on_page)])
    doc.multiBuild(story)
    return out_path
